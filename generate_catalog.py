import json
import re
import statistics
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "IST-layout-source.xlsm"
OUTPUT = ROOT / "dist" / "catalog.js"


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


def model_name(sheet_name):
    return " ".join(re.sub(r"^\s*Layout\s*", "", sheet_name, flags=re.I).split())


def category(name):
    value = normalize(name)
    if value == "IQC": return "KIỂM TRA ĐẦU VÀO"
    if "OQC" in value: return "KIỂM TRA ĐẦU RA"
    if "LASER WELD" in value: return "HÀN LASER"
    if any(x in value for x in ("VISION", "WLT", "VSWR", "WPT", "TEST", "CHECK")): return "TEST/KIỂM TRA"
    if "CLEAN" in value: return "LÀM SẠCH"
    if "METAL SHEET" in value: return "METAL SHEET"
    if "ROLLING" in value: return "ROLLING"
    if "AIR VENT" in value: return "AIRVENT"
    if "VINYL" in value: return "DÁN VINYL BẢO VỆ"
    return "DÁN & ÉP"


def number(value):
    if isinstance(value, bool) or value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


book = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
catalog, positions = {}, {}
for sheet_name in book.sheetnames:
    if not normalize(sheet_name).startswith("LAYOUT"):
        continue
    sheet = book[sheet_name]
    blocks = []
    ignored = {"MODEL", "TYPE", "LINE", "LEADER", "DKN", "SUPPORT", "TOTAL MP", "QC", "TOTAL", "PROCESS", "TT AVG", "MP", "NHAN VIEN", "ATTACH", "DINO", "WLT"}
    for row in range(1, sheet.max_row - 1):
        items = []
        for col in range(1, sheet.max_column + 1):
            raw = sheet.cell(row, col).value
            name = " ".join(raw.split()) if isinstance(raw, str) else ""
            if not name or normalize(name) in ignored:
                continue
            first, second = number(sheet.cell(row + 1, col).value), number(sheet.cell(row + 2, col).value)
            if first is None or second is None or not (0 <= first <= 300 and 0 <= second <= 300):
                continue
            stt = number(sheet.cell(row - 1, col).value) if row > 1 else None
            items.append({"name": name, "first": first, "second": second, "stt": int(stt) if stt and 0 < stt < 300 else None, "col": col})
        if len(items) >= 3:
            swap = statistics.median(x["first"] for x in items) <= 3 < statistics.median(x["second"] for x in items)
            for item in items:
                takt, mp = (item["second"], item["first"]) if swap else (item["first"], item["second"])
                item["time"], item["mp"] = round(takt, 3), int(mp) if 0 <= mp <= 30 else 0
            blocks.append(items)
    seen, extracted = set(), []
    for block_index, items in enumerate(blocks):
        for column_index, item in enumerate(sorted(items, key=lambda x: x["col"])):
            key = (item["stt"], normalize(item["name"]))
            if key in seen:
                continue
            seen.add(key)
            item.update({"block": block_index, "columnIndex": column_index})
            extracted.append(item)
    for index, item in enumerate(extracted):
        if not item["stt"]: item["stt"] = index + 1
    extracted.sort(key=lambda x: (x["stt"], x["block"], x["col"]))
    model = model_name(sheet_name)
    rows, layout_items = [], []
    for index, item in enumerate(extracted):
        rows.append({"name": item["name"], "category": category(item["name"]), "skill25": "", "skill40": "", "time": item["time"], "mp": item["mp"], "stt": item["stt"]})
        layout_items.append({"index": index, "name": item["name"], "stt": item["stt"], "x": item["columnIndex"] * 78, "y": 60 + item["block"] * 94, "time": item["time"], "mp": item["mp"]})
    if rows:
        catalog[model] = rows
        positions[model] = {"sheet": sheet_name, "matched": len(rows), "items": layout_items}

payload = "window.LOB_CATALOG=" + json.dumps(catalog, ensure_ascii=False, separators=(",", ":")) + ";\n"
payload += "window.LOB_POSITIONS=" + json.dumps(positions, ensure_ascii=False, separators=(",", ":")) + ";\n"
OUTPUT.write_text(payload, encoding="utf-8")
print(f"Generated {len(catalog)} models and {sum(map(len, catalog.values()))} processes directly from Layout sheets")
